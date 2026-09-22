import os
from pathlib import Path

import pymupdf
from PIL import Image


TEXT_EXTENSIONS = {".md", ".txt", ".csv", ".xml"}
OFFICE_EXTENSIONS = {".docx", ".xlsx"}
PDF_EXTENSIONS = {".pdf"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png"}

SUPPORTED_EXTENSIONS = (
    TEXT_EXTENSIONS
    | OFFICE_EXTENSIONS
    | PDF_EXTENSIONS
    | IMAGE_EXTENSIONS
)

# Resolucion de rasterizado de paginas PDF antes de enviarlas al modelo multimodal
PDF_RENDER_DPI = int(
    os.getenv("PDF_RENDER_DPI", "150")
)

# Lado maximo de la imagen enviada a Voyage (controla el consumo de tokens)
IMAGE_MAX_SIDE = int(
    os.getenv("IMAGE_MAX_SIDE", "1568")
)


def _downscale(image: Image.Image) -> Image.Image:
    if max(image.size) <= IMAGE_MAX_SIDE:
        return image

    ratio = IMAGE_MAX_SIDE / max(image.size)

    new_size = (
        int(image.width * ratio),
        int(image.height * ratio)
    )

    return image.resize(new_size, Image.LANCZOS)


def load_markdown(path: Path):
    text = path.read_text(encoding="utf-8")

    return [
        {
            "modality": "text",
            "text": text,
            "image": None,
            "page": None
        }
    ]


def load_pdf(path: Path):
    """Cada pagina se devuelve como bloque multimodal: texto + imagen renderizada."""
    document = pymupdf.open(path)

    pages = []

    for page_number, page in enumerate(document, start=1):
        text = page.get_text()

        pixmap = page.get_pixmap(dpi=PDF_RENDER_DPI)

        image = Image.frombytes(
            "RGB",
            (pixmap.width, pixmap.height),
            pixmap.samples
        )

        pages.append({
            "modality": "multimodal",
            "text": text.strip(),
            "image": _downscale(image),
            "page": page_number
        })

    document.close()

    return pages


def load_image(path: Path):
    image = Image.open(path).convert("RGB")

    return [
        {
            "modality": "image",
            "text": path.stem.replace("_", " "),
            "image": _downscale(image),
            "page": None
        }
    ]


def load_office(path: Path):
    text = ""
    extension = path.suffix.lower()
    
    if extension == ".docx":
        import docx
        doc = docx.Document(path)
        text = "\n".join([para.text for para in doc.paragraphs])
    elif extension == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"--- Sheet: {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True):
                # Filter out completely empty rows
                if not all(cell is None for cell in row):
                    lines.append(", ".join([str(cell) if cell is not None else "" for cell in row]))
        text = "\n".join(lines)
        
    return [
        {
            "modality": "text",
            "text": text,
            "image": None,
            "page": None
        }
    ]


def load_document(path: Path):

    extension = path.suffix.lower()

    if extension in PDF_EXTENSIONS:
        return load_pdf(path)

    if extension in TEXT_EXTENSIONS:
        return load_markdown(path)
        
    if extension in OFFICE_EXTENSIONS:
        return load_office(path)

    if extension in IMAGE_EXTENSIONS:
        return load_image(path)

    raise ValueError(
        f"Formato no soportado: {extension}"
    )
